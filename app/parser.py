from io import BytesIO
from pathlib import Path

from pypdf import PdfReader
from docx import Document


def extract_pdf_text(file_bytes: bytes) -> str:
    """
    Extract text from a PDF file.

    Parameters:
        file_bytes: PDF file content in bytes.

    Returns:
        Extracted text as a string.
    """

    reader = PdfReader(BytesIO(file_bytes))

    pages = []

    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)

    return "\n".join(pages)


def extract_docx_text(file_bytes: bytes) -> str:
    """
    Extract text from a DOCX file.

    Parameters:
        file_bytes: DOCX file content in bytes.

    Returns:
        Extracted text as a string.
    """

    document = Document(BytesIO(file_bytes))

    paragraphs = []

    # Extract normal paragraphs
    for paragraph in document.paragraphs:
        paragraphs.append(paragraph.text)

    # Extract text from tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                row_text.append(cell.text)

            paragraphs.append(" ".join(row_text))

    return "\n".join(paragraphs)


def extract_txt_text(file_bytes: bytes) -> str:
    """
    Extract text from a TXT file.

    Parameters:
        file_bytes: TXT file content in bytes.

    Returns:
        Extracted text as a string.
    """

    return file_bytes.decode("utf-8", errors="ignore")


def extract_resume_text(uploaded_file) -> str:
    """
    Detect the uploaded resume format and extract its text.

    Supported formats:
        PDF
        DOCX
        TXT
    """

    extension = Path(uploaded_file.name).suffix.lower()

    file_bytes = uploaded_file.getvalue()

    if extension == ".pdf":
        return extract_pdf_text(file_bytes)

    elif extension == ".docx":
        return extract_docx_text(file_bytes)

    elif extension == ".txt":
        return extract_txt_text(file_bytes)

    else:
        raise ValueError(
            "Unsupported file type. Please upload PDF, DOCX, or TXT."
        )